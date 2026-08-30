"""Dataset access, spreadsheet previews, and the seeded suite draw."""

import hashlib
import json
import random
from pathlib import Path

import openpyxl

from . import config, scorer


def load_dataset() -> list[dict]:
    path = config.DATASET_DIR / "dataset.json"
    if not path.exists():
        raise SystemExit(
            f"Dataset not found at {path}. Download and extract "
            "spreadsheetbench_verified_400.tar.gz into ~/.cache/skill-evolution-mvp "
            "(see IMPLEMENTATION.md)."
        )
    return json.loads(path.read_text())


def task_files(task: dict, test_case: int) -> tuple:
    """(input_path, answer_path). Verified-400 naming: {tc}_{id}_init.xlsx /
    {tc}_{id}_golden.xlsx (the 912 set uses _input/_answer; 395 of 400 tasks
    here have exactly one test case)."""
    d = config.DATASET_DIR / "spreadsheet" / str(task["id"])
    return (d / f"{test_case}_{task['id']}_init.xlsx",
            d / f"{test_case}_{task['id']}_golden.xlsx")


def test_cases(task: dict) -> list[int]:
    """Test-case numbers that actually exist for this task."""
    d = config.DATASET_DIR / "spreadsheet" / str(task["id"])
    out = []
    for tc in (1, 2, 3):
        if (d / f"{tc}_{task['id']}_init.xlsx").exists() and \
           (d / f"{tc}_{task['id']}_golden.xlsx").exists():
            out.append(tc)
    return out


def spreadsheet_preview(path, max_rows: int = 5, max_cols: int = 10,
                        max_chars: int = 2000) -> str:
    """First few rows of every sheet, tab-separated (the paper's
    'spreadsheet_content: The first few rows')."""
    try:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except Exception as e:
        return f"(could not read spreadsheet: {e})"
    parts = []
    for name in wb.sheetnames:
        ws = wb[name]
        parts.append(f"[Sheet: {name}] ({ws.max_row} rows x {ws.max_column} cols)")
        for i, row in enumerate(ws.iter_rows(max_row=max_rows, max_col=max_cols,
                                             values_only=True)):
            parts.append("\t".join("" if v is None else str(v)[:40] for v in row))
            if i + 1 >= max_rows:
                break
    wb.close()
    text = "\n".join(parts)
    return text[:max_chars] + ("\n..." if len(text) > max_chars else "")


def checker_selftest(task: dict) -> dict:
    """The suite-draw filter: on every available test case, answer-vs-answer
    must pass and input-vs-answer must fail (well-formed AND requires work)."""
    tcs = test_cases(task)
    if not tcs:
        return {"eligible": False, "reason": "no canonical init/golden files"}
    ok_answer, ok_input = [], []
    for tc in tcs:
        inp, ans = task_files(task, tc)
        a, _ = scorer.compare(ans, ans, task)
        b, _ = scorer.compare(ans, inp, task)
        ok_answer.append(a)
        ok_input.append(b)
    if not all(ok_answer):
        return {"eligible": False, "reason": "answer file fails its own check"}
    if any(ok_input):
        return {"eligible": False, "reason": "input already passes on some test case"}
    return {"eligible": True, "reason": ""}


def _published_splits() -> dict:
    """SkillOpt's published SpreadsheetBench id split (80/40/280) — the split
    WikiSkill reports strictly matching, over this exact Verified-400 tarball."""
    base = Path.home() / ".cache" / "skill-evolution-mvp" / "skillopt_splits"
    out = {"manifest": json.loads((base / "split_manifest.json").read_text())}
    for s in ("train", "val", "test"):
        items = json.loads((base / s / "items.json").read_text())
        out[s] = [str(x["id"] if isinstance(x, dict) else x) for x in items]
    return out


def draw_suite() -> dict:
    """Seeded, filtered draw of 30/15/100 + smoke, each NESTED inside the
    corresponding published SkillOpt split (distributional comparability at
    reduced cost). Deterministic."""
    dataset = load_dataset()
    by_id = {str(t["id"]): t for t in dataset}
    published = _published_splits()
    rng = random.Random(config.SUITE_SEED)

    splits, excluded = {}, []

    def draw_from(pool_ids: list, n: int, skip: set) -> list:
        pool = [i for i in pool_ids if i in by_id and i not in skip]
        rng.shuffle(pool)
        chosen = []
        for tid in pool:
            st = checker_selftest(by_id[tid])
            if st["eligible"]:
                chosen.append(tid)
            else:
                excluded.append({"id": tid, "reason": st["reason"]})
            if len(chosen) >= n:
                return chosen
        raise SystemExit(f"pool exhausted: got {len(chosen)} of {n}")

    splits["train"] = draw_from(published["train"], config.SPLIT_SIZES["train"], set())
    splits["val"] = draw_from(published["val"], config.SPLIT_SIZES["val"], set())
    splits["test"] = draw_from(published["test"], config.SPLIT_SIZES["test"], set())
    smoke = draw_from(published["train"], config.SMOKE_TASKS, set(splits["train"]))

    manifest = {
        "seed": config.SUITE_SEED,
        "dataset_tarball_sha256": config.DATASET_TARBALL_SHA256,
        "upstream_commit": config.UPSTREAM_COMMIT,
        "nested_within": {
            "source": "microsoft/SkillOpt data/spreadsheetbench_id_split",
            "counts": published["manifest"].get("counts"),
            "source_revision": published["manifest"].get("source_revision"),
        },
        "filter": "answer==answer passes AND input!=answer on every available test case",
        "splits": splits,
        "smoke": smoke,
        "excluded": excluded,
        "candidates_scanned": sum(len(v) for v in splits.values()) + len(smoke) + len(excluded),
    }
    blob = json.dumps(manifest, sort_keys=True).encode()
    manifest["manifest_sha256"] = hashlib.sha256(blob).hexdigest()
    config.TASKS_DIR.mkdir(exist_ok=True)
    out = config.TASKS_DIR / "manifest.json"
    out.write_text(json.dumps(manifest, indent=2))
    return manifest


def load_manifest() -> dict:
    return json.loads((config.TASKS_DIR / "manifest.json").read_text())


def tasks_by_id(ids: list) -> list[dict]:
    idx = {str(t["id"]): t for t in load_dataset()}
    return [idx[str(i)] for i in ids]
